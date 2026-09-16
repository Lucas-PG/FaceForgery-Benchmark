"""Gera a apresentação em pt-BR sem carregar imagens, pesos ou modelos do benchmark.

Uso: python presentation/build.py
As primitivas vetoriais são compartilhadas entre PowerPoint, PDF e HTML.
O notebook apenas reexpressa as médias publicadas em dados.json.
"""
from __future__ import annotations

import hashlib
import html
import json
import math
import os
import platform
import subprocess
from decimal import Decimal
from pathlib import Path

from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

ROOT = Path(__file__).resolve().parent
W, H, SCALE = 1600, 900, .6
BG, INK, MUTED = '#F7F9FC', '#17283F', '#526278'
TEAL, BLUE, LINE, WHITE, PALE = '#087F7B', '#315FA4', '#DDE4EC', '#FFFFFF', '#E7F3F1'
SAND = '#FFF1E5'
SLIDES: list[dict] = []


def configure_fonts() -> None:
    candidates = [
        Path('/usr/share/fonts/truetype/liberation2'),
        Path('/usr/share/fonts/truetype/liberation'),
        Path(os.environ.get('PRESENTATION_FONT_DIR', '/missing-fonts')),
    ]
    for directory in candidates:
        regular = directory / 'LiberationSans-Regular.ttf'
        bold = directory / 'LiberationSans-Bold.ttf'
        if regular.is_file() and bold.is_file():
            pdfmetrics.registerFont(TTFont('Deck', str(regular)))
            pdfmetrics.registerFont(TTFont('Deck-Bold', str(bold)))
            return
    raise RuntimeError('Instale fonts-liberation ou informe PRESENTATION_FONT_DIR com LiberationSans-Regular.ttf e LiberationSans-Bold.ttf. Fontes não são copiadas para a entrega.')


def num(value, digits=3) -> str:
    return f'{value:.{digits}f}'.replace('.', ',')


def delta(a, b) -> float:
    return float(Decimal(str(a)) - Decimal(str(b)))


def rect(s, x, y, w, h, fill=WHITE, radius=0, stroke=None):
    s['items'].append(dict(kind='rect', x=x, y=y, w=w, h=h, fill=fill, radius=radius, stroke=stroke))


def line(s, x1, y1, x2, y2, color=LINE, width=2):
    s['items'].append(dict(kind='line', x=x1, y=y1, x2=x2, y2=y2, color=color, width=width))


def dot(s, x, y, size=14, color=TEAL, square=False):
    s['items'].append(dict(kind='dot', x=x-size/2, y=y-size/2, w=size, h=size, color=color, square=square))


def text(s, value, x, y, w=1420, size=31, color=INK, bold=False, leading=1.24):
    font = 'Deck-Bold' if bold else 'Deck'
    lines = []
    for para in str(value).split('\n'):
        current = ''
        for word in para.split():
            if pdfmetrics.stringWidth(word, font, size) > w-8:
                raise ValueError(f'Palavra sem espaço excede a caixa: {word}')
            trial = (current + ' ' + word).strip()
            if current and pdfmetrics.stringWidth(trial, font, size) > w-8:
                lines.append(current); current = word
            else:
                current = trial
        lines.append(current)
    height = len(lines) * size * leading + 8
    if x < 0 or y < 0 or x+w > W or y+height > H:
        raise ValueError(f'Texto fora do slide: {value}')
    s['items'].append(dict(kind='text', x=x, y=y, w=w, h=height, lines=lines,
                           size=size, color=color, bold=bold, leading=leading))
    return height


def slide(title, tag='RESULTADOS REPORTADOS', source='A · Estudo anterior; B · seção V-A.', notes=''):
    s = dict(title=title, tag=tag, source=source, notes=notes, items=[])
    SLIDES.append(s)
    rect(s, 0, 0, W, H, BG)
    rect(s, 80, 64, 9, 23, TEAL, radius=4)
    text(s, tag, 107, 59, 1370, size=18, bold=True, color=TEAL)
    text(s, title, 80, 115, 1440, size=54, bold=True)
    line(s, 80, 811, 1520, 811)
    text(s, 'FONTES: '+source, 80, 833, 1320, size=17, color=MUTED)
    text(s, f'{len(SLIDES):02d}', 1460, 828, 60, size=23, bold=True, color=MUTED)
    return s


def paragraph_card(s, x, y, w, h, label, headline, body, fill=WHITE):
    rect(s, x, y, w, h, fill, radius=20)
    text(s, label, x+28, y+25, w-56, size=19, color=TEAL, bold=True)
    text(s, headline, x+28, y+72, w-56, size=39, bold=True)
    text(s, body, x+28, y+144, w-56, size=28, color=MUTED)


def export_pptx():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W/120), Inches(H/120)
    prs.core_properties.title = 'Detecção de falsificações faciais — resultados e explicabilidade'
    prs.core_properties.subject = 'Apresentação pt-BR; resultados publicados e protocolos pré-execução'
    prs.core_properties.author = 'Equipe de pesquisa — PUCPR'
    def pos(v): return Inches(v/120)
    def color(v): return RGBColor.from_string(v.lstrip('#'))
    for spec in SLIDES:
        out = prs.slides.add_slide(prs.slide_layouts[6])
        for item in spec['items']:
            k = item['kind']
            if k == 'text':
                box = out.shapes.add_textbox(pos(item['x']), pos(item['y']), pos(item['w']), pos(item['h']))
                tf = box.text_frame
                tf.clear(); tf.word_wrap = False
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.vertical_anchor = MSO_ANCHOR.TOP
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                for index, row in enumerate(item['lines']):
                    p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
                    p.text = row
                    p.font.name = 'Arial'
                    p.font.size = Pt(item['size']*SCALE)
                    p.font.bold = item['bold']
                    p.font.color.rgb = color(item['color'])
                    p.space_before = p.space_after = Pt(0)
                    p.line_spacing = Pt(item['size']*item['leading']*SCALE)
            elif k == 'line':
                sh = out.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, pos(item['x']), pos(item['y']), pos(item['x2']), pos(item['y2']))
                sh.line.color.rgb = color(item['color']); sh.line.width = Pt(item['width']*SCALE)
            else:
                kind = (MSO_SHAPE.RECTANGLE if item.get('square') else MSO_SHAPE.OVAL) if k == 'dot' else (MSO_SHAPE.ROUNDED_RECTANGLE if item.get('radius') else MSO_SHAPE.RECTANGLE)
                sh = out.shapes.add_shape(kind, pos(item['x']), pos(item['y']), pos(item['w']), pos(item['h']))
                sh.fill.solid(); sh.fill.fore_color.rgb = color(item.get('fill', item.get('color')))
                if item.get('stroke'):
                    sh.line.color.rgb = color(item['stroke']); sh.line.width = Pt(1)
                else: sh.line.fill.background()
                if k == 'rect' and item.get('radius'):
                    sh.adjustments[0] = min(.2, item['radius']/min(item['w'], item['h']))
        out.notes_slide.notes_text_frame.text = spec['notes']+'\n\nFontes: '+spec['source']+'\n'+SOURCE_INDEX
    prs.save(ROOT/'apresentacao.pptx')


def export_pdf():
    c = canvas.Canvas(str(ROOT/'apresentacao.pdf'), pagesize=(W*SCALE,H*SCALE), invariant=1)
    c.setTitle('Detecção de falsificações faciais — resultados e explicabilidade')
    c.setAuthor('Equipe de pesquisa — PUCPR')
    c.setSubject('pt-BR | Resultados reportados; protocolos de explicabilidade ainda não executados')
    for spec in SLIDES:
        c.saveState(); c.scale(SCALE,SCALE)
        for i in spec['items']:
            if i['kind'] == 'text':
                c.setFont('Deck-Bold' if i['bold'] else 'Deck',i['size']); c.setFillColor(HexColor(i['color']))
                for j,row in enumerate(i['lines']):
                    c.drawString(i['x'], H-i['y']-i['size']-j*i['size']*i['leading'], row)
            elif i['kind'] == 'line':
                c.setStrokeColor(HexColor(i['color'])); c.setLineWidth(i['width'])
                c.line(i['x'],H-i['y'],i['x2'],H-i['y2'])
            else:
                c.setFillColor(HexColor(i.get('fill',i.get('color'))))
                stroke=int(bool(i.get('stroke')))
                if stroke: c.setStrokeColor(HexColor(i['stroke']))
                x,y,w,h=i['x'],H-i['y']-i['h'],i['w'],i['h']
                if i['kind']=='dot' and not i.get('square'):
                    c.ellipse(x,y,x+w,y+h,stroke=0,fill=1)
                else: c.roundRect(x,y,w,h,i.get('radius',0),stroke=stroke,fill=1)
        c.restoreState(); c.showPage()
    c.save()


def export_html():
    sections=[]
    for idx,s in enumerate(SLIDES):
        parts=[f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="{html.escape(s["title"],quote=True)}"><title>{html.escape(s["title"])}</title>']
        for i in s['items']:
            if i['kind']=='text':
                weight='700' if i['bold'] else '400'
                for j,row in enumerate(i['lines']):
                    y=i['y']+i['size']+j*i['size']*i['leading']
                    parts.append(f'<text x="{i["x"]}" y="{y}" fill="{i["color"]}" font-size="{i["size"]}" font-weight="{weight}">{html.escape(row)}</text>')
            elif i['kind']=='line':
                parts.append(f'<line x1="{i["x"]}" y1="{i["y"]}" x2="{i["x2"]}" y2="{i["y2"]}" stroke="{i["color"]}" stroke-width="{i["width"]}"/>')
            elif i['kind']=='dot' and not i.get('square'):
                parts.append(f'<circle cx="{i["x"]+i["w"]/2}" cy="{i["y"]+i["h"]/2}" r="{i["w"]/2}" fill="{i["color"]}"/>')
            else:
                parts.append(f'<rect x="{i["x"]}" y="{i["y"]}" width="{i["w"]}" height="{i["h"]}" rx="{i.get("radius",0)}" fill="{i.get("fill",i.get("color"))}" stroke="{i.get("stroke") or "none"}"/>')
        parts.append('</svg>')
        sections.append(f'<section class="slide" id="slide-{idx+1}" aria-label="Slide {idx+1}">'+''.join(parts)+'</section>')
    document='''<!doctype html><html lang="pt-BR"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Resultados e explicabilidade — PUCPR</title><style>
*{box-sizing:border-box}body{margin:0;background:#17283f;font-family:Arial,sans-serif;color:white}main{height:calc(100vh - 62px);display:flex;align-items:center;justify-content:center}.slide{display:none;width:min(100vw,calc((100vh - 62px)*16/9));aspect-ratio:16/9}.slide.active{display:block}svg{width:100%;height:100%;display:block}nav{height:62px;display:flex;align-items:center;justify-content:center;gap:18px}button{font:inherit;background:#fff;color:#17283f;border:0;border-radius:8px;padding:10px 16px;cursor:pointer}button:focus-visible{outline:3px solid #58d5c0;outline-offset:3px}button:disabled{opacity:.5;cursor:default}#counter{min-width:72px;text-align:center}
@media print{@page{size:16in 9in;margin:0}body{background:white}main{display:block;height:auto}.slide,.slide.active{display:block;width:16in;height:9in;break-after:page}.slide:last-child{break-after:auto}nav{display:none}}
</style><main>'''+''.join(sections)+'''</main><nav aria-label="Navegação dos slides"><button id="prev" aria-label="Slide anterior">Anterior</button><span id="counter" aria-live="polite"></span><button id="next" aria-label="Próximo slide">Próximo</button><button id="full">Tela cheia</button></nav><script>
const slides=[...document.querySelectorAll('.slide')];let index=Math.max(0,Math.min(slides.length-1,(parseInt(location.hash.slice(1))||1)-1));
function show(){slides.forEach((s,i)=>{s.classList.toggle('active',i===index);s.setAttribute('aria-hidden',String(i!==index))});document.getElementById('counter').textContent=`${index+1} / ${slides.length}`;document.getElementById('prev').disabled=index===0;document.getElementById('next').disabled=index===slides.length-1;history.replaceState(null,'','#'+(index+1))}
function move(n){index=Math.max(0,Math.min(slides.length-1,index+n));show()}
document.getElementById('prev').onclick=()=>move(-1);document.getElementById('next').onclick=()=>move(1);document.getElementById('full').onclick=()=>{if(!document.fullscreenElement)document.documentElement.requestFullscreen?.();else document.exitFullscreen?.()};
document.addEventListener('keydown',e=>{if(e.target.tagName==='BUTTON'&&[' ','Enter'].includes(e.key))return;if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();move(1)}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();move(-1)}if(e.key==='Home'){index=0;show()}if(e.key==='End'){index=slides.length-1;show()}});show();</script></html>'''
    (ROOT/'apresentacao.html').write_text(document,encoding='utf-8')


def export_notebook(data):
    import nbformat
    from nbclient import NotebookClient
    cells=[]
    def md(value): cells.append(nbformat.v4.new_markdown_cell(value))
    def code(value): cells.append(nbformat.v4.new_code_cell(value))
    md('# Resultados do benchmark e explicabilidade\n\n**Material de apoio em português brasileiro.** Este notebook usa exclusivamente `dados.json`, cujos valores foram transcritos do estudo anterior e conferidos contra a edição de seis páginas. Não carrega pesos, não acessa imagens e não executa novos experimentos. ROC-AUC não é porcentagem de acerto.\n\nOs arquivos `apresentacao.pptx`, `apresentacao.pdf` e `apresentacao.html` contêm o mesmo roteiro visual. As notas completas estão em `roteiro.md`.')
    code("from pathlib import Path\nfrom decimal import Decimal\nimport json\nimport pandas as pd\nimport matplotlib.pyplot as plt\nfrom matplotlib.ticker import FuncFormatter\nfrom IPython.display import display\n\nfolder = Path.cwd() if (Path.cwd() / 'dados.json').is_file() else Path.cwd() / 'presentation'\nassert (folder / 'dados.json').is_file(), 'Abra o notebook na pasta presentation ou na raiz do repositório.'\ndata = json.loads((folder / 'dados.json').read_text(encoding='utf-8'))\nprint(data['proveniencia'])")
    md('## 1. Teste limpo versus degradado\n\nAs médias são do benchmark anterior: cinco famílias treinadas do zero e a configuração DINOv3 congelada. Elas não caracterizam os checkpoints fine-tuned da extensão. A queda é a diferença aritmética das médias arredondadas. Não é uma estimativa de significância.')
    code("rgb = pd.DataFrame(data['rgb'])\nrgb['queda_auc'] = [float(Decimal(str(a)) - Decimal(str(b))) for a, b in zip(rgb.limpo, rgb.degradado)]\ndisplay(rgb.round(3))\nassert len(rgb) == 6\nassert rgb.loc[rgb.modelo.eq('Xception'), 'queda_auc'].iloc[0] == 0.275\nassert rgb.loc[rgb.modelo.eq('DINOv3'), 'queda_auc'].iloc[0] == 0.083")
    code("ordered = rgb.sort_values('limpo', ascending=False).reset_index(drop=True)\nfig, ax = plt.subplots(figsize=(11, 5.8))\nfor i, row in ordered.iterrows():\n    ax.plot([row.limpo, row.degradado], [i, i], linewidth=1, alpha=0.4)\nax.scatter(ordered.limpo, range(len(ordered)), label='Teste limpo', marker='o', s=65)\nax.scatter(ordered.degradado, range(len(ordered)), label='Teste degradado', marker='s', s=65)\nax.set_yticks(range(len(ordered)), ordered.modelo)\nax.invert_yaxis()\nax.set_xlim(0.5, 1.0)\nax.xaxis.set_major_formatter(FuncFormatter(lambda x, _: f'{x:.2f}'.replace('.', ',')))\nax.set_xlabel('ROC-AUC — médias reportadas; sem intervalos inferidos')\nax.set_title('A degradação muda o ranking')\nax.legend(frameon=False)\nax.grid(axis='x', alpha=0.2)\nfig.tight_layout()\nplt.show()\nplt.close(fig)")
    md('**Leitura:** Xception lidera o teste limpo (0,884); DINOv3 tem a maior média no degradado (0,726) entre estas seis configurações. A diferença de regime de pré-treinamento impede atribuir essa comparação apenas à arquitetura. A pequena queda do ViT (0,016) ocorre a partir de uma média limpa já baixa (0,624).')
    md('## 2. RGB versus entradas híbridas no teste degradado\n\nOs três exemplos são contrastes descritivos reportados. Não representam todos os modos possíveis e não demonstram causalidade nem generalização para outros datasets.')
    code("hybrid = pd.DataFrame(data['hibridos'])\nhybrid['ganho_auc'] = [float(Decimal(str(a)) - Decimal(str(b))) for a, b in zip(hybrid.hibrido_degradado, hybrid.rgb_degradado)]\ndisplay(hybrid.round(3))\nassert hybrid.ganho_auc.tolist() == [0.041, 0.013, 0.018]")
    code("fig, ax = plt.subplots(figsize=(10, 4.5))\npositions = list(range(len(hybrid)))\nax.scatter(hybrid.rgb_degradado, positions, label='RGB', marker='o', s=65)\nax.scatter(hybrid.hibrido_degradado, positions, label='Híbrido reportado', marker='s', s=65)\nax.set_yticks(positions, hybrid.modelo)\nax.invert_yaxis()\nax.set_xlim(0.5, 0.75)\nax.xaxis.set_major_formatter(FuncFormatter(lambda x, _: f'{x:.2f}'.replace('.', ',')))\nax.set_title('Frequência como complemento: três contrastes reportados')\nax.set_xlabel('ROC-AUC no teste degradado')\nax.legend(frameon=False)\nax.grid(axis='x', alpha=0.2)\nfig.tight_layout()\nplt.show()\nplt.close(fig)")
    md('## 3. As três perguntas de explicabilidade\n\n**Tarefa 1:** congelar 64 imagens, com 16 VP, FN, VN e FP definidos pelo modelo de referência; reutilizar as mesmas identidades em todos os modelos. A distribuição 16×4 não é garantida para os demais modelos.\n\n**Tarefa 2:** selecionar somente na validação um par competente RGB–frequência da mesma arquitetura, regime e semente. Comparar cada atribuição no domínio nativo: pixels em RGB, coeficientes no espectro. Não há heatmaps novos neste notebook.\n\n**Tarefa 3:** calcular a interseção das falhas do conjunto declarado, sem descartar predições faltantes silenciosamente, e comparar com controles pareados por rótulo. Não chamar essas falhas de universais.\n\nA semente principal é 42. Falsa é a classe positiva. O alvo das atribuições específicas de classe é a diferença entre o logit falso e o real.')
    md('## 4. Métodos e cuidados\n\n'+ '\n\n'.join('**'+k+':** '+', '.join(v)+'.' for k,v in data['metodos'].items())+'\n\nO suporte depende da arquitetura. Attention Rollout é agnóstico à classe. Concordância visual não substitui checagens numéricas, sensibilidade ao baseline ou controles de randomização e perturbação.')
    md('## 5. Estado da evidência\n\n'+ '\n'.join('- '+v for v in data['limites'])+'\n\nNenhuma taxa de falhas compartilhadas ou conclusão demográfica foi estimada neste material. A execução futura fornece esses resultados; o código e a literatura não os predeterminam.')
    md('## Fontes\n\n'+ '\n\n'.join(f'**[{r["id"]}] {r["titulo"]}** — {r["localizacao"]}. [Consultar]({r["url"]})' for r in data['fontes'])+'\n\nReferências completas da pesquisa: `paper/explicability/references.bib`.')
    for n,cell in enumerate(cells):
        cell['id']=hashlib.sha256((str(n)+cell.source).encode()).hexdigest()[:12]
        cell['metadata']['slideshow']={'slide_type': 'slide' if cell.cell_type=='markdown' else 'fragment'}
    nb=nbformat.v4.new_notebook(cells=cells,metadata={'kernelspec':{'name':'python3','display_name':'Python 3','language':'python'},'language_info':{'name':'python','version':platform.python_version()}})
    NotebookClient(nb,timeout=60,kernel_name='python3',resources={'metadata':{'path':str(ROOT)}}).execute()
    nbformat.validate(nb)
    assert not any(o.output_type=='error' for c in nb.cells if c.cell_type=='code' for o in c.outputs)
    nbformat.write(nb,ROOT/'apresentacao.ipynb')


def validate_and_preview(data):
    import pymupdf
    import nbformat
    from PIL import Image
    pdf=pymupdf.open(ROOT/'apresentacao.pdf')
    ppt=Presentation(ROOT/'apresentacao.pptx')
    assert len(pdf)==len(ppt.slides)==len(SLIDES)==15
    for n,page in enumerate(pdf):
        assert len(page.get_text().strip()) > 100, f'Página vazia: {n+1}'
        for word in page.get_text('words'):
            x0,y0,x1,y1=word[:4]
            assert x0>=0 and y0>=0 and x1<=page.rect.width+.5 and y1<=page.rect.height+.5, f'Texto fora da página {n+1}'
    for n,s in enumerate(SLIDES):
        boxes=[i for i in s['items'] if i['kind']=='text']
        for aidx,a in enumerate(boxes):
            # Compare actual rendered line extent, not spare textbox width.
            aw=max(pdfmetrics.stringWidth(t,'Deck-Bold' if a['bold'] else 'Deck',a['size']) for t in a['lines'])
            for b in boxes[aidx+1:]:
                bw=max(pdfmetrics.stringWidth(t,'Deck-Bold' if b['bold'] else 'Deck',b['size']) for t in b['lines'])
                overlap_x=min(a['x']+aw,b['x']+bw)-max(a['x'],b['x'])
                overlap_y=min(a['y']+a['h']-8,b['y']+b['h']-8)-max(a['y'],b['y'])
                if overlap_x>2 and overlap_y>2:
                    raise ValueError(f'Textos sobrepostos no slide {n+1}: {a["lines"]} / {b["lines"]}')
    overview=Image.new('RGB',(1200,5*225),(227,233,241))
    for idx,page in enumerate(pdf):
        pix=page.get_pixmap(matrix=pymupdf.Matrix(400/page.rect.width,400/page.rect.width),alpha=False)
        image=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
        overview.paste(image,((idx%3)*400,(idx//3)*225))
    overview.save(ROOT/'visao-geral.png')
    # Large preview pages are CI-only; fonts and model/data files are never packaged.
    previews=Path('/tmp/presentation-previews'); previews.mkdir(exist_ok=True)
    for idx,page in enumerate(pdf):
        page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False).save(previews/f'slide-{idx+1:02}.png')
    files=['apresentacao.pptx','apresentacao.pdf','apresentacao.html','apresentacao.ipynb','roteiro.md','visao-geral.png']
    report={'idioma':'pt-BR','slides':15,'formato':'16:9','fonte_dos_dados':data['source_commit'],
            'status':'Resultados reportados no estudo anterior; nenhum experimento novo de XAI.',
            'verificacoes':['15 slides no PowerPoint e no PDF','notebook executado sem erro; dados publicados apenas','textos dentro das páginas','ausência de sobreposição entre caixas de texto','AUC e diferenças derivadas verificadas','HTML independente de rede'],
            'sha256':{p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in files},
            'source_sha256':{p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in ['dados.json','conteudo.py','build.py','requirements.txt']}}
    (ROOT/'build.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))


def main():
    global SOURCE_INDEX
    configure_fonts()
    data=json.loads((ROOT/'dados.json').read_text(encoding='utf-8'))
    assert sum(map(len,data['metodos'].values()))==12
    assert all(0<=r['degradado']<=r['limpo']<=1 for r in data['rgb'])
    SOURCE_INDEX='\n'.join(f'[{r["id"]}] {r["titulo"]}: {r["url"]}' for r in data['fontes'])
    from conteudo import compose
    compose(__import__(__name__),data)
    assert len(SLIDES)==15
    export_pptx(); export_pdf(); export_html(); export_notebook(data)
    script=['# Roteiro de apresentação — pt-BR','',data['proveniencia'],'']
    for n,s in enumerate(SLIDES,1):
        script += [f'## {n:02}. {s["title"]}','',s['notes'],'','Fontes: '+s['source'],'']
    script += ['## Referências','',SOURCE_INDEX]
    (ROOT/'roteiro.md').write_text('\n'.join(script)+'\n',encoding='utf-8')
    validate_and_preview(data)


if __name__=='__main__':
    main()
